"""Selected workbook categories must never overwrite one another silently."""
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from PyFLASH.batch import Batch
from PyFLASH.conditions import condition, conditionList


def batch_with_literal_measurements(folder):
    groups = conditionList([condition(name, name, '#112233', 'Group')
                            for name in ('B', 'A', 'Absent')])
    summary = pd.DataFrame({'Condition': ['A','B','A','B'],
        'AnimalName': ['first','second','third','fourth'],
        'IBA1_Count': [2.,5.,7.,11.], 'CustomScore': [2.,5.,7.,11.]}, index=[11,2,17,5])
    behavior = summary[['Condition','AnimalName']].copy()
    behavior['Period'] = [102.,105.,107.,111.]
    batch = object.__new__(Batch)
    batch.condition_list, batch.conditions = groups, groups.conditions
    batch.factor, batch.factorDict = groups.factor, groups.factorDict
    batch.summaries, batch.summary = {'SCN': summary}, summary
    batch.data = {'IBA1': SimpleNamespace(df=summary), 'Behaviour': SimpleNamespace(df=behavior)}
    batch.experiment_list, batch.export_path = [], str(folder)
    return batch


@pytest.mark.parametrize('method', ['export_excel','export_all_excel'])
@pytest.mark.parametrize('behavior_name, extra_name', [
    ('Shared','Shared'), (' Shared.xlsx ','Shared'), ('Shared.XLSX','Shared.xlsx')])
@pytest.mark.parametrize('previous', [False,True])
def test_colliding_names_are_refused_before_any_selected_output_is_written(
        tmp_path, method, behavior_name, extra_name, previous):
    batch = batch_with_literal_measurements(tmp_path)
    snapshots = {name: obj.df.copy(deep=True) for name,obj in batch.data.items()}
    existing = {'Shared.xlsx': b'accepted previous workbook',
                'Shared_RegexFilters.txt': b'accepted previous filter report'} if previous else {}
    for name, contents in existing.items():
        (tmp_path / name).write_bytes(contents)
    with pytest.raises(ValueError, match='overwrite one another'):
        getattr(batch, method)(tmp_path, if_summary=False, if_extended=False,
            behaviour_save_name=behavior_name, unregistered_summary_save_name=extra_name)
    assert {p.name:p.read_bytes() for p in tmp_path.iterdir()} == existing
    for name, frame in snapshots.items():
        pd.testing.assert_frame_equal(batch.data[name].df, frame, check_exact=True)


@pytest.mark.parametrize('method', ['export_excel','export_all_excel'])
def test_distinct_requested_outputs_preserve_both_literal_measurement_tables(tmp_path, method):
    batch = batch_with_literal_measurements(tmp_path)
    outputs = getattr(batch, method)(tmp_path, if_summary=False, if_extended=False,
        behaviour_include='^Period$', unregistered_summary_include='^CustomScore$',
        behaviour_save_name='SelectedBehavior', unregistered_summary_save_name='SelectedExtras')
    assert [Path(p).name for p in outputs] == [
        'SelectedBehavior.xlsx','SelectedBehavior_RegexFilters.txt',
        'SelectedExtras.xlsx','SelectedExtras_RegexFilters.txt']
    tables = []
    for path in outputs:
        if not str(path).endswith('.xlsx'):
            continue
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                if rows and rows[0] == ['B','A','Absent']:
                    tables.append(rows[1:3])
        finally:
            book.close()
    assert tables == [[[105,102,None],[111,107,None]], [[5,2,None],[11,7,None]]]


@pytest.mark.parametrize('method', ['export_excel','export_all_excel'])
def test_disabled_category_does_not_reserve_its_unused_filename(tmp_path, method):
    batch = batch_with_literal_measurements(tmp_path)
    outputs = getattr(batch, method)(tmp_path, if_summary=False, if_extended=False,
        behaviour=False, behaviour_save_name='Shared', unregistered_summary_save_name='Shared')
    assert [Path(path).name for path in outputs] == ['Shared.xlsx','Shared_RegexFilters.txt']


@pytest.mark.parametrize('method', ['export_excel','export_all_excel'])
def test_region_exports_with_colliding_names_are_refused_before_root_outputs(tmp_path, method):
    batch = batch_with_literal_measurements(tmp_path)
    with pytest.raises(ValueError, match='overwrite one another'):
        getattr(batch, method)(tmp_path, if_summary_save_name='Shared', if_extended_save_name='Shared')
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('method', ['export_excel','export_all_excel'])
def test_two_region_keys_cannot_overwrite_a_shared_lowercase_folder(tmp_path, method):
    batch = batch_with_literal_measurements(tmp_path)
    batch.summaries['scn'] = batch.summary.copy()
    with pytest.raises(ValueError, match='overwrite one another'):
        getattr(batch, method)(tmp_path)
    assert list(tmp_path.iterdir()) == []
