"""Every public batch export parameter must preserve its literal transport."""
import pytest
import shutil
from pathlib import Path
from reproducibility import excel_orchestration_contract as reference


def test_public_export_parameter_surface_is_completely_classified():
    assert len(reference.check_parameter_coverage()) == 17


@pytest.mark.parametrize('row', reference.cases(), ids=lambda row: row['key'])
def test_actual_workbook_selection_paths_and_complete_measurement_tables(row, tmp_path):
    observed = reference.observe(row, tmp_path)
    assert observed['inputs_unchanged']


@pytest.mark.parametrize('name', ['client-reference.json', 'declared-client-recipe.py', 'source-manifest.json'])
def test_changed_original_evidence_is_refused(name, tmp_path, monkeypatch):
    archive = tmp_path / 'archive'
    shutil.copytree(reference.ARCHIVE, archive)
    path = archive / name
    path.write_bytes(path.read_bytes() + b' ')
    monkeypatch.setattr(reference, 'ARCHIVE', archive)
    with pytest.raises(ValueError, match='Original batch export reference'):
        reference.check_identity()


def _one_summary_case():
    return next(row for row in reference.cases() if row['kind'] == 'ordinary'
                and row['path_mode'] == 'absolute' and row['roi_layout'] == 'one'
                and row['flags'] == dict(behaviour=False, if_extended=False,
                                        if_summary=True, unregistered_summary=False))


def test_changed_exported_measurement_fails_the_frozen_contract(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from PyFLASH.batch import Batch
    original = Batch.export_excel

    def corrupt_measurement(batch, *args, **kwargs):
        paths = original(batch, *args, **kwargs)
        path = next(Path(path) for path in paths if path.endswith('.xlsx'))
        book = load_workbook(path)
        try:
            sheet = next(sheet for sheet in book if [cell.value for cell in sheet[1]] == ['B', 'A', 'Absent'])
            sheet.cell(row=2, column=1, value=999)
            book.save(path)
        finally:
            book.close()
        return paths

    monkeypatch.setattr(Batch, 'export_excel', corrupt_measurement)
    with pytest.raises(AssertionError):
        reference.observe(_one_summary_case(), tmp_path)


def test_changed_client_input_is_refused_before_any_export(tmp_path, monkeypatch):
    original = reference.literal_batch

    def changed_input(*args, **kwargs):
        batch = original(*args, **kwargs)
        batch.summary.loc[11, 'IBA1_Count'] = 999
        return batch

    monkeypatch.setattr(reference, 'literal_batch', changed_input)
    with pytest.raises(AssertionError, match='Original client input tables changed'):
        reference.observe(_one_summary_case(), tmp_path)
    assert not list(tmp_path.rglob('*.xlsx'))
