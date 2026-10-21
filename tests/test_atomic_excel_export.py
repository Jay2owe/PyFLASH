"""A failed workbook attempt preserves the complete accepted file and report."""
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from PyFLASH import _persistence
from PyFLASH.batch import Batch
from PyFLASH.conditions import condition, conditionList
from PyFLASH.export import _literal_excel_writer


def previous_book(path):
    book = Workbook()
    book.active.title = 'Accepted'
    book.active.append(['subject', 'value'])
    book.active.append(['original', 7])
    book.save(path)
    book.close()
    original = path.read_bytes()
    inputs = path.parent.parent / 'prior-inputs'
    inputs.mkdir(exist_ok=True)
    (inputs / (path.parent.name + '-' + path.name)).write_bytes(original)
    return original


@pytest.mark.parametrize('previous', [False, True])
@pytest.mark.parametrize('stage', ['body', 'close', 'replacement'])
def test_failed_workbook_keeps_previous_complete_bytes(tmp_path, monkeypatch, previous, stage):
    target = tmp_path / 'accepted.xlsx'
    before = previous_book(target) if previous else None
    error = OSError('declared workbook ' + stage + ' failure')
    if stage == 'replacement':
        def replace(*args):
            raise error
        monkeypatch.setattr(_persistence, 'atomic_replace', replace)
    elif stage == 'close':
        import xlsxwriter
        original = xlsxwriter.Workbook.close
        def close(book):
            original(book)
            raise error
        monkeypatch.setattr(xlsxwriter.Workbook, 'close', close)
    with pytest.raises(OSError) as caught:
        with _literal_excel_writer(target) as writer:
            pd.DataFrame({'subject': ['new'], 'value': [19]}).to_excel(writer, index=False)
            if stage == 'body':
                raise error
    assert caught.value is error
    assert target.read_bytes() == before if previous else not target.exists()
    assert sorted(path.name for path in tmp_path.iterdir()) == (['accepted.xlsx'] if previous else [])


def test_successful_workbook_publishes_all_literal_cells(tmp_path):
    target = tmp_path / 'complete.xlsx'
    with _literal_excel_writer(target) as writer:
        pd.DataFrame({'subject': ['=SUM(A1:A2)'], 'value': [3]}).to_excel(
            writer, sheet_name='Data', index=False)
    book = load_workbook(target, data_only=False)
    try:
        assert book.sheetnames == ['Data']
        assert list(book['Data'].values) == [('subject', 'value'), ('=SUM(A1:A2)', 3)]
        assert book['Data']['A2'].data_type == 's'
    finally:
        book.close()
    assert [path.name for path in tmp_path.iterdir()] == ['complete.xlsx']


@pytest.mark.parametrize('method', ['export_behavior_summary_excel', 'export_excel', 'export_all_excel'])
@pytest.mark.parametrize('previous', [False, True])
@pytest.mark.parametrize('error_type', [OSError, ValueError])
def test_behavior_sheet_write_failure_is_not_reported_as_success(
        tmp_path, monkeypatch, method, previous, error_type):
    owner = conditionList([condition('B', 'B', '#112233', 'Group'),
                           condition('A', 'A', '#334455', 'Group')])
    frame = pd.DataFrame({'Condition': ['A', 'B'], 'AnimalName': ['one', 'two'],
                          'Period': [24., 25.]})
    original_frame = frame.copy(deep=True)
    batch = object.__new__(Batch)
    batch.condition_list, batch.conditions = owner, owner.conditions
    batch.factor, batch.factorDict = owner.factor, owner.factorDict
    batch.data = {'Behaviour': SimpleNamespace(df=frame)}
    batch.summaries, batch.experiment_list = {}, []
    batch.summary = pd.DataFrame()
    target = tmp_path / 'Behavior_Summary.xlsx'
    report = tmp_path / 'Behavior_Summary_RegexFilters.txt'
    before = previous_book(target) if previous else None
    if previous:
        report.write_bytes(b'accepted report\n')
    error = error_type('declared measurement sheet write failure')
    original = pd.DataFrame.to_excel
    def broken(table, writer, *args, **kwargs):
        result = original(table, writer, *args, **kwargs)
        if kwargs.get('sheet_name') != 'Conditions':
            raise error
        return result
    monkeypatch.setattr(pd.DataFrame, 'to_excel', broken)
    options = dict(include='^Period$') if method == 'export_behavior_summary_excel' else dict(
        if_summary=False, if_extended=False, unregistered_summary=False, behaviour_include='^Period$')
    with pytest.raises(error_type) as caught:
        getattr(batch, method)(save_path=tmp_path, **options)
    assert caught.value is error
    pd.testing.assert_frame_equal(frame, original_frame, check_exact=True)
    assert target.read_bytes() == before if previous else not target.exists()
    assert sorted(path.name for path in tmp_path.iterdir()) == (
        ['Behavior_Summary.xlsx', 'Behavior_Summary_RegexFilters.txt'] if previous else [])
    if previous:
        assert report.read_bytes() == b'accepted report\n'
