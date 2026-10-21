"""Literal region-selection workbook controls; other export options remain open."""
import ast
import hashlib
import json
from pathlib import Path

ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-region'
MANIFEST_SHA256 = '8805bd049c7faeb98c2f5933c45de456efe7090b46342b1da276ea9e224dd159'
BEFORE = 'excel-region-functional-before-20261007-v1'
CAPABILITY_PREFIXES = {'PyFLASH.batch.Batch.export_extended_data_excel': ('excel-region.',)}
COMPLETE_CAPABILITIES = frozenset()

CACHE_CARRIER_SHA256 = 'efdf97fa3ab6403499be0dd1f9b82ad3630c8dd2d1de19812cb62dd58be83eb9'


def check_identity():
    from reproducibility.evidence_identity import verify_manifest
    verify_manifest(ARCHIVE, 'source-manifest.json', MANIFEST_SHA256, cache_carrier_sha256=CACHE_CARRIER_SHA256)


def cases():
    return json.loads((ARCHIVE / BEFORE / 'frozen-before.json').read_text(encoding='utf-8'))['rows']

def production(row, output):
    import pandas as pd
    from openpyxl import load_workbook
    from PyFLASH.batch import Batch
    from PyFLASH.conditions import condition, conditionList
    from PyFLASH.markers import Antibody
    from types import SimpleNamespace
    filename = ARCHIVE / BEFORE / 'freeze_excel_roi_selection_20261007.py'
    parsed = ast.parse(filename.read_text(encoding='utf-8'))
    function = next(node for node in parsed.body
                    if isinstance(node, ast.FunctionDef) and node.name == 'production')
    namespace = {'Path': Path, 'SimpleNamespace': SimpleNamespace, 'pd': pd,
                 'load_workbook': load_workbook, 'Batch': Batch, 'Antibody': Antibody,
                 'condition': condition, 'conditionList': conditionList}
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(filename), 'exec'), namespace)
    return namespace['production'](row, Path(output) / row['key'])

def agrees(row, actual):
    expected = row['original']['expected_table']
    return (actual['error'] is None and actual['actual_table'] == expected
            and actual['metric_sheet_count'] == int(expected is not None)
            and actual['inputs_unchanged']
            and actual['output_names'] == ['literal-' + row['key'] + '.xlsx',
                                          'literal-' + row['key'] + '_RegexFilters.txt'])

def validate(output):
    import openpyxl
    from PyFLASH._persistence import write_json
    from reproducibility.normality_association_validation import comparison
    check_identity()
    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    rows, paired = [], []
    scope = ('Literal acquired-region metadata and selector controls through the actual '
             'extended-workbook export, including repeated rows and unchanged inputs. '
             'No complete export-options or microscopy scientific certificate.')
    for row in cases():
        actual = production(row, output)
        equal = agrees(row, actual)
        expected = row['original']['expected_table']
        rows.append(comparison(row['key'], 'Literal supplied-value contract and native openpyxl reader',
                               'excel-region', actual['actual_table'], expected,
                               equal, True, scope, 0, 0))
        paired.append({'key': row['key'], 'actual': actual, 'expected_table': expected,
                       'agrees': equal, 'scope': scope})
        write_json(output / 'excel-region-paired.json', {
            'reader': 'openpyxl', 'reader_version': openpyxl.__version__,
            'manifest_sha256': MANIFEST_SHA256, 'rows': paired})
    return rows
