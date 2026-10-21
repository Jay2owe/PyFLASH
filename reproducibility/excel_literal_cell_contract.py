"""Literal identities/text and unlisted groups must survive Excel export.

Native openpyxl writes and independently reads explicitly literal reference
cells; the actual PyFLASH files are written by XlsxWriter. This functional
schema agreement is not an independently fitted scientific estimate.
"""
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from reproducibility import excel_orchestration_contract as original
from reproducibility import excel_workbook_schema_contract as workbook_schema

ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-literal-cells' / 'v1'
MANIFEST_SHA256 = '391dcd33dc123de3f2b7248d128a068e4d081a82bfe53b823d7d8cf430240765'
COMPLETE_CAPABILITIES = frozenset()
CAPABILITY_PREFIXES = {'PyFLASH.batch.Batch.' + method: ('excel-literal-cells.' + method + '.',)
                       for method in ('export_excel', 'export_all_excel', 'export_IF_summary_excel',
                       'export_unregistered_summary_excel', 'export_extra_summary_excel',
                       'export_behavior_summary_excel', 'export_extended_data_excel')}


def check_identity():
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    if sha(ARCHIVE/'source-manifest.json') != MANIFEST_SHA256:
        raise ValueError('Original literal Excel cells manifest changed')
    for name, digest in json.loads((ARCHIVE/'source-manifest.json').read_text(encoding='utf-8')).items():
        if sha(ARCHIVE/name) != digest:
            raise ValueError('Original literal Excel cells evidence changed: ' + name)
    fixture = json.loads((ARCHIVE/'client-reference.json').read_text(encoding='utf-8'))
    original.check_identity()
    assert fixture['original_input_manifest_sha256'] == original.MANIFEST_SHA256
    return fixture


def cases():
    return check_identity()['rows']


def check_original_failures():
    archive = ARCHIVE.parent/'original-probes/v1'
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    if sha(archive/'source-manifest.json') != '65e6276cb42f1e865bc5606d0e99668dcc49125e0f2f8c95233bf7ef529cac0a':
        raise ValueError('Original literal Excel failure manifest changed')
    for name, digest in json.loads((archive/'source-manifest.json').read_text(encoding='utf-8')).items():
        if sha(archive/name) != digest:
            raise ValueError('Original literal Excel failure evidence changed: ' + name)
    receipt = json.loads((archive/'run.json').read_text(encoding='utf-8'))
    assert receipt['comparisons'] == 49 and receipt['disagreements'] == 27 and not receipt['source_changes']
    return receipt


def native_literal_reference(row, folder):
    from openpyxl import Workbook
    paths = []
    for index, expected in enumerate(row['expected_workbooks']):
        book = Workbook()
        book.remove(book.active)
        for schema in expected['worksheets']:
            sheet = book.create_sheet(schema['name'])
            for i, values in enumerate(schema['cells'], 1):
                for j, value in enumerate(values, 1):
                    cell = sheet.cell(i, j, value)
                    if isinstance(value, str):
                        # openpyxl's public cell type explicitly writes literal
                        # strings; native formula defaults are not our contract.
                        cell.data_type = 's'
            for merged in schema['merged_ranges']:
                sheet.merge_cells(merged)
        path = folder / f'native-literal-reference-{index}.xlsx'
        try:
            book.save(path)
        finally:
            book.close()
        actual = workbook_schema.read_workbook(path)
        normalized = [dict(name=sheet['name'], cells=sheet['cells'], merged_ranges=sheet['merged_ranges'])
                      for sheet in actual]
        assert normalized == expected['worksheets'], 'Native literal writer changed its declared client cells'
        assert not any(sheet['formula_cells'] for sheet in actual)
        paths.append(dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    return paths


def literal_batch(folder, row):
    base = next(case for case in original.cases() if case['key'] == row['base_input_key'])
    batch = original.literal_batch(folder, base)
    profile = row['profile']
    if profile == 'formula-group-name':
        batch.condition_list[0].name = '=SUM(1,2)'
        for frame in [batch.summary, *batch.summaries.values(), *[obj.df for obj in batch.data.values()]]:
            frame['Condition'] = frame['Condition'].replace({'B': '=SUM(1,2)'})
    if profile == 'formula-explanation':
        batch.conditions[0].factor_explanation = '=SUM(1,2)'
    if profile == 'formula-subject-name':
        for frame in [batch.summary, *batch.summaries.values(), *[obj.df for obj in batch.data.values()]]:
            frame['AnimalName'] = frame['AnimalName'].replace({'first': '=SUM(1,2)'})
    if profile == 'formula-extra-text':
        batch.summary['CustomScore'] = ['=SUM(1,2)', '=1+1', '@literal', '+literal']
    if profile in ('unknown-raw-group', 'missing-raw-group'):
        batch.data['IBA1'].df.loc[0, 'Condition'] = 'Unlisted' if profile == 'unknown-raw-group' else None
    return batch


def observe(row, output):
    import pandas as pd
    folder = (Path(output) / row['key']).resolve()
    folder.mkdir(parents=True, exist_ok=False)
    natives = native_literal_reference(row, folder)
    batch = literal_batch(folder, row)
    frames = [batch.summary, *batch.summaries.values(), *[obj.df for obj in batch.data.values()]]
    snapshots = [(frame, frame.copy(deep=True)) for frame in frames]
    groups = deepcopy([condition.__dict__ for condition in batch.conditions])
    input_tables = original._input_tables(batch)
    paths, error = [], None
    try:
        paths = getattr(batch, row['method'])(save_path=folder/'absolute', **row['options'])
    except Exception as exc:
        error = type(exc).__name__ + ': ' + str(exc)
    for frame, snapshot in snapshots:
        pd.testing.assert_frame_equal(frame, snapshot, check_exact=True)
        assert frame.attrs == snapshot.attrs
    assert groups == [condition.__dict__ for condition in batch.conditions]
    expected_paths = [str(folder/item['request']['directory']/(item['request']['name']+ending))
                      for item in row['expected_workbooks'] for ending in ('.xlsx', '_RegexFilters.txt')]
    actual_paths = [str(Path(path).resolve()) for path in paths or []]
    workbooks = []
    for expected in row['expected_workbooks']:
        item = expected['request']
        path = folder/item['directory']/(item['name']+'.xlsx')
        if path.is_file():
            workbooks.append(dict(request=item, worksheets=workbook_schema.read_workbook(path),
                regex_report=path.with_name(item['name']+'_RegexFilters.txt').read_text(encoding='utf-8')
                if path.with_name(item['name']+'_RegexFilters.txt').exists() else None))
    observed = dict(error=error, outputs=actual_paths, inputs_unchanged=True, input_tables=input_tables,
                    workbooks=workbooks, native_literal_references=natives)
    observed['disagreements'] = workbook_schema.evaluate(row, observed)
    if error != row['expected_error']:
        observed['disagreements'].append('unexpected execution error')
    if actual_paths != expected_paths:
        observed['disagreements'].append('output paths/order')
    return observed


def required_keys(capability):
    method = capability.rsplit('.', 1)[-1]
    return {row['key'] for row in cases() if row['method'] == method}


def validate(output):
    import traceback
    from PyFLASH._persistence import write_json
    from reproducibility.normality_association_validation import comparison
    paired, comparisons = [], []
    for row in cases():
        observed, error = None, None
        try:
            observed = observe(row, output)
        except Exception:
            error = traceback.format_exc()
        agrees = error is None and not observed['disagreements']
        comparisons.append(comparison(row['key'], 'predeclared-literal-client/native-openpyxl-literal-writer-and-reader',
            'functional-literal-data-and-identity-transport/no-additional-estimator-vote',
            {'literal_cells_and_complete_schema_match': agrees}, {'literal_cells_and_complete_schema_match': True},
            agrees, True, 'exact-input-identities-literal-strings-all-sheets-and-reports', 0., 0.))
        paired.append(dict(case=row, observed=observed, error=error, agrees=agrees))
    write_json(Path(output)/'excel-literal-cells-paired.json', dict(schema_version=1,
        complete_contract=False, scope=__doc__, rows=paired))
    return comparisons
