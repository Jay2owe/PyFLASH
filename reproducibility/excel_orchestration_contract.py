"""Whole public batch-export routing over literal supplied measurement tables.

The subject is selection, option forwarding, effective output paths and complete
exported measurement arrays, not fitting or microscopy preprocessing. Native
openpyxl independently reads the actual XlsxWriter files. This is a functional
transport contract, with no extra independent estimator vote.
"""
from itertools import product
from pathlib import Path
import inspect
import os
import hashlib
import json
import math

ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-orchestration' / 'v1'
MANIFEST_SHA256 = '78f3174841cd0a38dfb6f2d9179a4d1e9aee68c358cddfcc5e1f746eeb5cf5e3'

METHODS = ('export_excel', 'export_all_excel')
FLAGS = ('behaviour', 'if_extended', 'if_summary', 'unregistered_summary')
FILTERS = tuple(flag + suffix for flag in FLAGS for suffix in ('_include', '_exclude'))
NAMES = tuple(flag + '_save_name' for flag in FLAGS)
CAPABILITY_PREFIXES = {'PyFLASH.batch.Batch.' + method: ('excel-orchestration.' + method + '.',)
                       for method in METHODS}
COMPLETE_CAPABILITIES = frozenset()  # Admit only after the complete original matrix qualifies.
DEFAULT_NAMES = {'behaviour': 'Behavior_Summary', 'if_extended': 'IF_Extended',
                 'if_summary': 'IF_Summary', 'unregistered_summary': 'Extra_Summary'}
METRIC_NAMES = {'behaviour': 'Period', 'if_extended': 'IBA1_Volume',
                'if_summary': 'IBA1_Count', 'unregistered_summary': 'CustomScore'}


def _declared_cases():
    rows = []
    for method, flags, path, rois, behavior in product(METHODS, product((False, True), repeat=4),
            ('absolute', 'relative', 'default'), ('one', 'two'), (False, True)):
        key = 'excel-orchestration.' + method + '.flags-' + ''.join(str(int(v)) for v in flags)
        rows.append(dict(key=key + '.' + path + '.' + rois + '.' + str(behavior), method=method,
            flags=dict(zip(FLAGS, flags)), path_mode=path, roi_layout=rois,
            behavior_present=behavior, options={}, kind='ordinary'))
    for method in METHODS:
        for parameter in FILTERS:
            for value in ('keep', 'drop', 'invalid'):
                flag = parameter.rsplit('_', 1)[0]
                matched = '^' + METRIC_NAMES[flag] + '$'
                unmatched = '^never-a-supplied-metric$'
                pattern = (matched if value == 'keep' else unmatched) if parameter.endswith('_include') else (
                    unmatched if value == 'keep' else matched)
                if value == 'invalid':
                    pattern = '['
                rows.append(dict(key='excel-orchestration.' + method + '.' + parameter + '.' + value,
                    method=method, flags={name: name == flag for name in FLAGS},
                    path_mode='absolute', roi_layout='two', behavior_present=True,
                    options={parameter: pattern}, kind='invalid-regex' if value == 'invalid' else 'filter'))
        for value in ('custom', 'none', 'extension', 'whitespace', 'empty', 'blank'):
            options = {parameter: ('Chosen-' + flag if value == 'custom' else None if value == 'none'
                else DEFAULT_NAMES[flag] + '.XLSX' if value == 'extension'
                else '' if value == 'empty' else '   ' if value == 'blank'
                else '  ' + DEFAULT_NAMES[flag] + '  ')
                for flag, parameter in zip(FLAGS, NAMES)}
            rows.append(dict(key='excel-orchestration.' + method + '.names-' + value,
                method=method, flags={name: True for name in FLAGS}, path_mode='absolute',
                roi_layout='two', behavior_present=True, options=options, kind='ordinary'))
        for kind in ('collision', 'unknown-option', 'empty-input', 'missing-input', 'legacy-summary', 'create-paths'):
            options = ({'behaviour_save_name': 'Shared', 'unregistered_summary_save_name': 'Shared'}
                       if kind == 'collision' else {'unrecognized_option': True} if kind == 'unknown-option' else {})
            rows.append(dict(key='excel-orchestration.' + method + '.' + kind, method=method,
                flags={name: True for name in FLAGS}, path_mode='default' if kind == 'create-paths' else 'absolute',
                roi_layout='one' if kind == 'legacy-summary' else 'two',
                behavior_present=True, options=options, kind=kind))
        for kind in ('defaults', 'disabled-options', 'region-name-collision', 'region-directory-collision'):
            options = ({**{name: '[' for name in FILTERS}, **{name: 'Shared' for name in NAMES}}
                if kind == 'disabled-options' else
                {'if_summary_save_name': 'Shared', 'if_extended_save_name': 'Shared'}
                if kind == 'region-name-collision' else {})
            rows.append(dict(key='excel-orchestration.' + method + '.' + kind, method=method,
                flags={name: kind != 'disabled-options' for name in FLAGS}, path_mode='absolute',
                roi_layout='two', behavior_present=True, options=options, kind=kind))
        for path in ('absolute', 'relative', 'default'):
            rows.append(dict(key='excel-orchestration.' + method + '.relative-base.' + path,
                method=method, flags={name: True for name in FLAGS}, path_mode=path,
                roi_layout='two', behavior_present=True, options={}, kind='relative-base'))
    return rows


def check_identity():
    """Authenticate original literal declarations and retained failed probes."""
    manifest_path = ARCHIVE / 'source-manifest.json'
    if hashlib.sha256(manifest_path.read_bytes()).hexdigest() != MANIFEST_SHA256:
        raise ValueError('Original batch export reference manifest changed')
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    for name, digest in manifest.items():
        if hashlib.sha256((ARCHIVE / name).read_bytes()).hexdigest() != digest:
            raise ValueError('Original batch export reference evidence changed: ' + name)
    fixture = json.loads((ARCHIVE / 'client-reference.json').read_text(encoding='utf-8'))
    if [item['case'] for item in fixture['rows']] != _declared_cases():
        raise ValueError('Declared batch export cases differ from frozen original requests')
    return fixture


def cases():
    return [dict(item['case'], _reference={key: value for key, value in item.items() if key != 'case'})
            for item in check_identity()['rows']]


def _input_tables(batch):
    def encode(value):
        if isinstance(value, float) and not math.isfinite(value):
            return {'nonfinite': 'nan' if math.isnan(value) else 'inf' if value > 0 else '-inf'}
        if isinstance(value, dict):
            return {key: encode(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [encode(item) for item in value]
        return value
    table = lambda frame: encode(frame.to_dict(orient='split'))
    return dict(summary=table(batch.summary),
                summaries={key: table(frame) for key, frame in batch.summaries.items()},
                data={key: table(marker.df) for key, marker in batch.data.items()})


def check_parameter_coverage():
    from PyFLASH.batch import Batch
    expected = {'save_path', *FLAGS, *FILTERS, *NAMES}
    actual = set(inspect.signature(Batch.export_excel).parameters) - {'self'}
    assert actual == expected, (actual, expected)
    assert inspect.signature(Batch.export_all_excel).parameters['kwargs'].kind is inspect.Parameter.VAR_KEYWORD
    return expected


def literal_batch(folder, row):
    import pandas as pd
    from types import SimpleNamespace
    from PyFLASH.batch import Batch
    from PyFLASH.conditions import condition, conditionList
    from PyFLASH.markers import Antibody
    groups = conditionList([condition(name, name, '#112233', 'Group') for name in ('B', 'A', 'Absent')])
    summary = pd.DataFrame({'Condition': ['A', 'B', 'A', 'B'], 'Group': ['A', 'B', 'A', 'B'],
        'AnimalName': ['first', 'second', 'third', 'fourth'],
        'IBA1_Count': [2., 5., 7., 11.], 'CustomScore': [12., 15., 17., 21.]}, index=[11,2,17,5])
    summaries = {'SCN': summary}
    if row['roi_layout'] == 'two':
        other = summary.copy(deep=True)
        other['IBA1_Count'] = [32., 35., 37., 41.]
        summaries['OC'] = other
    raw = pd.DataFrame({'Condition': ['A','B','A','B','A','B','A','B'],
        'AnimalName': ['first','second','third','fourth'] * 2,
        'ROI': ['SCN'] * 4 + ['OC'] * 4,
        'IBA1_Volume': [42.,45.,47.,51.,62.,65.,67.,71.]})
    marker = object.__new__(Antibody)
    marker.name, marker.df = 'IBA1', raw
    experiment = SimpleNamespace(data={'IBA1': marker}, summary=summary, filePath=str(folder))
    batch = object.__new__(Batch)
    batch.condition_list, batch.conditions = groups, groups.conditions
    batch.factor, batch.factorDict = groups.factor, groups.factorDict
    batch.summary, batch.summaries = summary, summaries
    batch.data = {'IBA1': marker}
    if row['behavior_present']:
        behavior = summary[['Condition','AnimalName']].copy()
        behavior['Period'] = [102.,105.,107.,111.]
        batch.data['Behaviour'] = SimpleNamespace(df=behavior)
    batch.experiment_list, batch.export_path = [experiment], str(folder / 'base')
    if row['kind'] == 'missing-input':
        for frame in batch.summaries.values():
            frame.loc[11, 'IBA1_Count'] = float('nan')
        batch.summary.loc[11, 'CustomScore'] = float('nan')
        batch.data['Behaviour'].df.loc[11, 'Period'] = float('nan')
        raw.loc[[0,4], 'IBA1_Volume'] = float('nan')
    if row['kind'] == 'empty-input':
        batch.summaries = {name: frame.iloc[:0].copy() for name,frame in batch.summaries.items()}
        batch.summary = batch.summaries['SCN']
        marker.df = raw.iloc[:0].copy()
        batch.data['Behaviour'].df = batch.data['Behaviour'].df.iloc[:0].copy()
        experiment.summary = batch.summary
    if row['kind'] == 'legacy-summary':
        batch.summaries = {}
    if row['kind'] == 'create-paths':
        del batch.export_path
        batch.filePath = str(folder / 'auto')
    if row['kind'] == 'relative-base':
        batch.export_path = 'base'
    if row['kind'] == 'region-directory-collision':
        batch.summaries['scn'] = summary.copy(deep=True)
    return batch


def expected_outputs(row):
    """Literal client request interpretation, without calling exporter helpers."""
    if '_reference' in row:
        return [{key: value for key, value in item.items() if key != 'measurements'}
                for item in row['_reference']['outputs']]
    result = []
    root = ('auto/Exports' if row['kind'] == 'create-paths' else
            'base/nested' if row['path_mode'] == 'relative' else 'base' if row['path_mode'] == 'default' else 'absolute')
    regions = ('SCN', 'OC') if row['roi_layout'] == 'two' else ('SCN',)
    for region in regions:
        for flag in ('if_summary', 'if_extended'):
            if row['flags'][flag]:
                result.append((flag, root + '/' + region.lower(), region))
    for flag in ('behaviour', 'unregistered_summary'):
        if row['flags'][flag] and (flag != 'behaviour' or row['behavior_present']):
            result.append((flag, root, None))
    outputs = []
    for flag, directory, region in result:
        name = row['options'].get(flag + '_save_name', DEFAULT_NAMES[flag])
        name = DEFAULT_NAMES[flag] if name is None else name.strip()
        if not name:
            name = DEFAULT_NAMES[flag]
        if name.lower().endswith('.xlsx'):
            name = name[:-5]
        outputs.append(dict(kind=flag, directory=directory, region=region, name=name))
    return outputs


def expected_measurements(row, item):
    """Exact supplied cells and explicitly selected metric absence."""
    if '_reference' in row:
        return next(output['measurements'] for output in row['_reference']['outputs']
                    if {key: value for key, value in output.items() if key != 'measurements'} == item)
    flag, region = item['kind'], item['region']
    include = row['options'].get(flag + '_include')
    exclude = row['options'].get(flag + '_exclude')
    if include == '^never-a-supplied-metric$' or exclude == '^' + METRIC_NAMES[flag] + '$':
        return None
    if row['kind'] == 'empty-input':
        return None if flag == 'if_extended' else [['B','A','Absent']]
    if flag == 'if_extended':
        values = ([42,47,45,51] if region == 'SCN' else [62,67,65,71])
        if row['kind'] == 'missing-input':
            values[0] = None
        return [['Condition','Animal ID','ROI','Volume (\u00b5m\u00b3)'],
                ['B','second',region,values[2]], ['B','fourth',region,values[3]],
                ['A','first',region,values[0]], ['A','third',region,values[1]]]
    body = ([[5,2,None],[11,7,None]] if flag == 'if_summary' and region == 'SCN' else
            [[35,32,None],[41,37,None]] if flag == 'if_summary' else
            [[105,102,None],[111,107,None]] if flag == 'behaviour' else
            [[15,12,None],[21,17,None]])
    if row['kind'] == 'missing-input':
        body = [[body[0][0],body[1][1],None],[body[1][0],None,None]]
    return [['B','A','Absent']] + body


def observe(row, output):
    import pandas as pd
    from openpyxl import load_workbook
    folder = (Path(output) / row['key']).resolve()
    folder.mkdir(parents=True, exist_ok=False)
    batch = literal_batch(folder, row)
    assert _input_tables(batch) == row['_reference']['input_tables'], 'Original client input tables changed'
    frames = [batch.summary, *batch.summaries.values(), *[obj.df for obj in batch.data.values()]]
    snapshots = [(frame, frame.copy(deep=True)) for frame in frames]
    supplied_path = folder / 'absolute' if row['path_mode'] == 'absolute' else 'nested' if row['path_mode'] == 'relative' else None
    options = {} if row['kind'] == 'defaults' else {**row['flags'], **row['options']}
    error = None
    outputs = []
    original_directory = Path.cwd()
    returned_paths = []
    try:
        if row['kind'] == 'relative-base':
            os.chdir(folder)
        outputs = getattr(batch, row['method'])(save_path=supplied_path, **options)
        returned_paths = list(outputs)
        outputs = [str(Path(path).resolve()) for path in outputs]
    except Exception as exc:
        error = type(exc).__name__ + ': ' + str(exc)
    finally:
        if row['kind'] == 'relative-base':
            os.chdir(original_directory)
    for frame, snapshot in snapshots:
        pd.testing.assert_frame_equal(frame, snapshot, check_exact=True)
        assert frame.attrs == snapshot.attrs
    error_prefix = row['_reference']['expected_error_prefix']
    if error_prefix is not None:
        assert error is not None and error.startswith(error_prefix), error
        assert outputs == []
        assert not list(folder.rglob('*.xlsx'))
        assert not list(folder.rglob('*_RegexFilters.txt'))
        return dict(error=error, outputs=[], inputs_unchanged=True)
    assert error is None, error
    expected = expected_outputs(row)
    expected_paths = [str(folder / item['directory'] / (item['name'] + ending))
                      for item in expected for ending in ('.xlsx','_RegexFilters.txt')]
    assert outputs == expected_paths, (outputs, expected_paths)
    assert all(Path(path).is_file() for path in outputs)
    measured = []
    for item in expected:
        path = folder / item['directory'] / (item['name'] + '.xlsx')
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheets = {sheet.title: [list(values) for values in sheet.iter_rows(values_only=True)] for sheet in book}
        finally:
            book.close()
        target = expected_measurements(row, item)
        tables = []
        for cells in worksheets.values():
            if cells and (cells[0] == ['B','A','Absent'] or cells[0] == ['Condition','Animal ID','ROI','Volume (\u00b5m\u00b3)']):
                length = len(target) if target is not None else len(cells)
                tables.append(cells[:length])
                # Empty measurements may still have a blank separator and the
                # declared prose description. Neither is a scientific row.
                # Extra numeric values or numeric strings remain disagreements.
                for remaining in cells[length:]:
                    assert all(value is None for value in remaining[1:])
                    value = remaining[0]
                    assert value is None or isinstance(value, str)
                    if value is not None:
                        try:
                            float(value)
                        except ValueError:
                            pass
                        else:
                            raise AssertionError('Unexpected additional numerical measurement')
        assert tables == ([] if target is None else [target]), (item, tables, target)
        measured.append(dict(request=item, worksheets=worksheets, expected_measurements=target))
    return dict(error=None, outputs=outputs, returned_paths=returned_paths,
                workbooks=measured, inputs_unchanged=True)


def required_keys(capability):
    method = capability.rsplit('.', 1)[-1]
    return {row['key'] for row in cases() if row['method'] == method}


def validate(output):
    import json
    import os
    from PyFLASH._persistence import write_json
    from reproducibility.normality_association_validation import comparison
    import traceback
    parameters = sorted(check_parameter_coverage())
    check_identity()
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    rows, paired = [], []
    # Each completed case is durable without rewriting all preceding answers.
    # Exclusive creation protects earlier failed attempts from being overwritten.
    with (output / 'excel-orchestration-paired.jsonl').open('x', encoding='utf-8') as journal:
        for row in cases():
            observed, error = None, None
            try:
                observed = observe(row, output)
            except Exception:
                error = traceback.format_exc()
            agrees = error is None
            rows.append(comparison(row['key'], 'literal-client-selected-output-schema/native-openpyxl-reader',
                'functional-complete-batch-export-routing/no-additional-estimator-vote',
                {'complete_requested_transport_matches': agrees}, {'complete_requested_transport_matches': True},
                agrees, True, 'exact-output-path-order-and-complete-supplied-measurement-arrays', 0.,0.))
            paired.append(dict(case=row, observed=observed, error=error, agrees=agrees))
            journal.write(json.dumps(paired[-1], sort_keys=True) + '\n')
            journal.flush()
            os.fsync(journal.fileno())
    write_json(output / 'excel-orchestration-paired.json', dict(schema_version=1,
        complete_contract=False, scope=__doc__, covered_parameters=parameters, rows=paired))
    return rows
