"""Literal spreadsheet row-order controls; whole exporter options remain open."""
import ast
import hashlib
import json
from pathlib import Path

ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-index'
TYPED_ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-typed-filters'
TYPED_MANIFEST_SHA256 = '585b53223a0fa9dedc81e0ec41cf2f1778e18bc29fd5ee0301f07891adb3275b'
MANIFEST_SHA256 = '63fb57b9debd675431c1b871eeb174a26859d9360689592a4cc85b516524a04f'
BEFORE = 'excel-index-functional-before-20261007-v1'
CAPABILITY_PREFIXES = {
    'PyFLASH.batch.Batch.' + name: ('excel-index.' + name + '.', 'excel-typed-filter.' + name + '.')
    for name in ('export_IF_summary_excel', 'export_unregistered_summary_excel',
                 'export_extra_summary_excel', 'export_behavior_summary_excel')
}
COMPLETE_CAPABILITIES = frozenset()


def check_identity():
    manifest = ARCHIVE / 'source-manifest.json'
    if hashlib.sha256(manifest.read_bytes()).hexdigest() != MANIFEST_SHA256:
        raise ValueError('Spreadsheet evidence manifest changed.')
    for name, digest in json.loads(manifest.read_text(encoding='utf-8')).items():
        if hashlib.sha256((ARCHIVE / name).read_bytes()).hexdigest() != digest:
            raise ValueError('Spreadsheet evidence changed: ' + name)
    typed_manifest = TYPED_ARCHIVE / 'source-manifest.json'
    if hashlib.sha256(typed_manifest.read_bytes()).hexdigest() != TYPED_MANIFEST_SHA256:
        raise ValueError('Typed spreadsheet evidence manifest changed.')
    for name, digest in json.loads(typed_manifest.read_text(encoding='utf-8')).items():
        if hashlib.sha256((TYPED_ARCHIVE / name).read_bytes()).hexdigest() != digest:
            raise ValueError('Typed spreadsheet evidence changed: ' + name)


def cases():
    return json.loads((ARCHIVE / BEFORE / 'frozen-before.json').read_text(encoding='utf-8'))['rows']


def typed_cases():
    return json.loads((TYPED_ARCHIVE / 'run.json').read_text(encoding='utf-8'))['rows']


def typed_domain_cases():
    """Explicit empty/refused client inputs, separate from original failed cases."""
    rows = []
    for row in typed_cases():
        if row['form'] not in ('numpy-array', 'pandas-series', 'pandas-index'):
            continue
        rows.append(dict(row, key=row['key'] + '.empty', form='empty-' + row['form']))
        rows.append(dict(row, key=row['key'] + '.invalid', literal_pattern='[', expected_refusal=True))
    return rows


def typed_production(row, output):
    """Replay unchanged literal inputs/reader with the named supplied container."""
    import numpy as np
    import pandas as pd
    from openpyxl import load_workbook
    from PyFLASH.batch import Batch
    from PyFLASH.conditions import condition, conditionList
    from types import SimpleNamespace
    original = next(value for value in cases() if value['method'] == row['method']
                    and value['index'] == [11, 2, 17, 5] and not value['missing']
                    and value['order'] == 'BA')
    pattern = row['literal_pattern']
    forms = {'string': lambda: pattern, 'list': lambda: [None, '', pattern],
             'tuple': lambda: (None, '', pattern), 'set': lambda: {None, '', pattern},
             'numpy-array': lambda: np.array(['', pattern, ''], dtype=object),
             'pandas-series': lambda: pd.Series([None, pattern, '']),
             'pandas-index': lambda: pd.Index([None, '', pattern])}
    form = row['form']
    if form.startswith('empty-'):
        supplied = {'numpy-array': lambda: np.array([], dtype=object),
                    'pandas-series': lambda: pd.Series([], dtype=object),
                    'pandas-index': lambda: pd.Index([], dtype=object)}[form[6:]]()
    else:
        supplied = forms[form]()
    actual_input = dict(original, key=row['key'], filter_target=row['target'],
                        supplied_filter=supplied)
    filename = TYPED_ARCHIVE / 'executed-input-and-reader-recipe.py'
    namespace = dict(Path=Path, SimpleNamespace=SimpleNamespace, pd=pd,
                     load_workbook=load_workbook, Batch=Batch,
                     condition=condition, conditionList=conditionList)
    exec(compile(filename.read_text(encoding='utf-8'), str(filename), 'exec'), namespace)
    return namespace['production'](actual_input, Path(output) / row['key'])


def typed_agrees(row, actual):
    if row.get('expected_refusal'):
        return (isinstance(actual['error'], str)
                and actual['error'].startswith('ValueError: Invalid ' + row['target'] + ' regex')
                and actual['workbook_cells'] == {} and actual['output_names'] == []
                and actual['inputs_unchanged'])
    return (actual['error'] is None and actual['metric_table_count'] == 1
            and actual['actual_metric_rows'] == row['expected_metric_rows']
            and actual['inputs_unchanged']
            and actual['output_names'] == ['literal-' + row['key'] + '.xlsx',
                                          'literal-' + row['key'] + '_RegexFilters.txt'])


def production(row, output):
    import pandas as pd
    from openpyxl import load_workbook
    from PyFLASH.batch import Batch
    from PyFLASH.conditions import condition, conditionList
    filename = ARCHIVE / BEFORE / 'freeze_excel_index_contracts_20261007.py'
    parsed = ast.parse(filename.read_text(encoding='utf-8'))
    function = next(node for node in parsed.body
                    if isinstance(node, ast.FunctionDef) and node.name == 'production')
    # Reuse the authenticated original input/reader recipe, without executing
    # its old output-directory launcher or importing a frozen production copy.
    namespace = {'Path': Path, 'SimpleNamespace': __import__('types').SimpleNamespace,
                 'pd': pd, 'load_workbook': load_workbook, 'Batch': Batch,
                 'condition': condition, 'conditionList': conditionList}
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(filename), 'exec'), namespace)
    return namespace['production'](row, Path(output) / row['key'])


def agrees(row, actual):
    return (actual['error'] is None and actual['metric_table_count'] == 1
            and actual['actual_metric_rows'] == row['original']['expected_metric_rows']
            and actual['inputs_unchanged']
            and actual['output_names'] == ['literal-' + row['key'] + '.xlsx',
                                          'literal-' + row['key'] + '_RegexFilters.txt'])


def validate(output):
    import openpyxl
    from PyFLASH._persistence import write_json
    from reproducibility.normality_association_validation import comparison
    check_identity()
    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    rows, paired = [], []
    scope = ('Literal spreadsheet source-row order, missing cells, declared group '
             'order, absent groups, output names and input ownership through actual '
             'public exports. Functional boundary controls; no complete exporter '
             'option or microscopy scientific certificate.')
    for row in cases():
        actual = production(row, output)
        equal = agrees(row, actual)
        expected = row['original']['expected_metric_rows']
        rows.append(comparison(row['key'], 'Literal supplied-value contract and native openpyxl reader',
                               'excel-index', actual['actual_metric_rows'], expected,
                               equal, True, scope, 0, 0))
        paired.append({'key': row['key'], 'actual': actual, 'expected_metric_rows': expected,
                       'agrees': equal, 'scope': scope})
        write_json(output / 'excel-index-paired.json', {
            'reader': 'openpyxl', 'reader_version': openpyxl.__version__,
            'manifest_sha256': MANIFEST_SHA256, 'rows': paired})
    typed_pairs = []
    for row in typed_cases() + typed_domain_cases():
        actual = typed_production(row, output)
        equal = typed_agrees(row, actual)
        refused = row.get('expected_refusal', False)
        observed = {'refused_before_workbook': equal} if refused else actual['actual_metric_rows']
        expected = {'refused_before_workbook': True} if refused else row['expected_metric_rows']
        rows.append(comparison(row['key'],
            'Literal supplied-value contract and native openpyxl reader',
            'excel-typed-filter', observed, expected,
            equal, True, scope + ' Equivalent include/exclude filter container transport.', 0, 0))
        typed_pairs.append({'key': row['key'], 'form': row['form'], 'target': row['target'],
                            'actual': actual, 'expected': expected, 'expected_refusal': refused,
                            'agrees': equal, 'scope': scope})
        write_json(output / 'excel-typed-filter-paired.json', {
            'reader': 'openpyxl', 'reader_version': openpyxl.__version__,
            'manifest_sha256': TYPED_MANIFEST_SHA256, 'complete_contract': False,
            'rows': typed_pairs})
    return rows
