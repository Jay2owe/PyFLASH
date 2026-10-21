"""Exact supporting workbook/report transport for every original batch request.

Native openpyxl reads the real XlsxWriter exports. Independent, predeclared
literal client schemas cover sheet order, all cells, descriptions, merged
ranges and every filter-report line. Existing numerical/routing controls stay
required. This does not certify untested exporter input domains or add an
independent scientific estimator.
"""
import hashlib
import json
from pathlib import Path

from reproducibility import excel_orchestration_contract as original

ARCHIVE = Path(__file__).with_name('reference-sources') / 'excel-workbook-schema' / 'v3'
MANIFEST_SHA256 = '3c666da16c73faac7d825da8ba0b458c8622d4331f2c6e0657553c9e9aa8186d'
COMPLETE_CAPABILITIES = frozenset()
CAPABILITY_PREFIXES = {'PyFLASH.batch.Batch.' + method: ('excel-workbook-schema.' + method + '.',)
                       for method in original.METHODS}


def check_identity():
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    if sha(ARCHIVE / 'source-manifest.json') != MANIFEST_SHA256:
        raise ValueError('Original complete workbook schema manifest changed')
    manifest = json.loads((ARCHIVE / 'source-manifest.json').read_text(encoding='utf-8'))
    for name, digest in manifest.items():
        if sha(ARCHIVE / name) != digest:
            raise ValueError('Original complete workbook schema evidence changed: ' + name)
    fixture = json.loads((ARCHIVE / 'client-reference.json').read_text(encoding='utf-8'))
    old = original.check_identity()
    if fixture['original_manifest_sha256'] != original.MANIFEST_SHA256:
        raise ValueError('Complete workbook schema original input identity changed')
    assert [(row['original_key'], row['method']) for row in fixture['rows']] == [
        (row['case']['key'], row['case']['method']) for row in old['rows']]
    return fixture


def cases():
    return check_identity()['rows']


def check_original_failures():
    archive = ARCHIVE.parent/'original-probes/v1'
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    if sha(archive/'source-manifest.json') != 'be59cc599c6df07537a1d698e24ba724555fd6b5dc66d9ad11740337e04719e4':
        raise ValueError('Original workbook schema failure manifest changed')
    for name, digest in json.loads((archive/'source-manifest.json').read_text(encoding='utf-8')).items():
        if sha(archive/name) != digest:
            raise ValueError('Original workbook schema failure evidence changed: ' + name)
    receipt = json.loads((archive/'run.json').read_text(encoding='utf-8'))
    assert receipt['comparisons'] == 470 and receipt['disagreements'] == 16 and not receipt['source_changes']
    return receipt


def check_empty_index_originals():
    archive = ARCHIVE.parent/'empty-index-original-probes/v1'
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    if sha(archive/'source-manifest.json') != '8e2297f61eb25724d8db98ef1651f219aeefb478d46dddbc4ad6c8664c313644':
        raise ValueError('Original empty-index failure manifest changed')
    for name,digest in json.loads((archive/'source-manifest.json').read_text(encoding='utf-8')).items():
        if sha(archive/name) != digest:
            raise ValueError('Original empty-index failure evidence changed: ' + name)
    rows = json.loads((archive/'excel-workbook-schema-paired.json').read_text(encoding='utf-8'))['rows']
    failed = [row for row in rows if not row['agrees']]
    assert len(failed) == 4 and all(row['observed']['disagreements'] == ['workbook 0, sheet 1: cells'] for row in failed)
    return failed


def read_workbook(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=False, data_only=False)
    try:
        worksheets = []
        for sheet in book:
            cells = [list(values) for values in sheet.iter_rows(values_only=True)]
            formulas = [cell.coordinate for row in sheet for cell in row if cell.data_type == 'f']
            worksheets.append(dict(name=sheet.title, cells=cells,
                merged_ranges=sorted(str(value) for value in sheet.merged_cells.ranges),
                formula_cells=formulas))
        return worksheets
    finally:
        book.close()


def evaluate(row, observed):
    expected = row['expected_workbooks']
    actual = observed['workbooks']
    disagreements = []
    if len(actual) != len(expected):
        disagreements.append('workbook count')
    for index, (got, want) in enumerate(zip(actual, expected)):
        if got['request'] != want['request']:
            disagreements.append(f'workbook {index}: request')
        sheets = [dict(name=sheet['name'], cells=sheet['cells'], merged_ranges=sheet['merged_ranges'])
                  for sheet in got['worksheets']]
        if sheets != want['worksheets']:
            for j, (actual_sheet, expected_sheet) in enumerate(zip(sheets, want['worksheets'])):
                for key in ('name', 'cells', 'merged_ranges'):
                    if actual_sheet[key] != expected_sheet[key]:
                        disagreements.append(f'workbook {index}, sheet {j}: {key}')
            if len(sheets) != len(want['worksheets']):
                disagreements.append(f'workbook {index}: worksheet count/order')
        if any(sheet['formula_cells'] for sheet in got['worksheets']):
            disagreements.append(f'workbook {index}: literal data became a formula')
        if got['regex_report'] != want['regex_report']:
            disagreements.append(f'workbook {index}: complete regex report')
    return disagreements


def observe(row, output):
    original_row = next(case for case in original.cases() if case['key'] == row['original_key'])
    # Existing source/input/path/measurement checks execute the actual public API.
    previous = original.observe(original_row, output)
    folder = (Path(output) / row['original_key']).resolve()
    workbooks = []
    for item in original.expected_outputs(original_row) if row['expected_error_prefix'] is None else []:
        path = folder / item['directory'] / (item['name'] + '.xlsx')
        workbooks.append(dict(request=item, worksheets=read_workbook(path),
            regex_report=path.with_name(item['name'] + '_RegexFilters.txt').read_text(encoding='utf-8')))
    observed = dict(error=previous['error'], outputs=previous['outputs'], workbooks=workbooks,
                    inputs_unchanged=previous['inputs_unchanged'])
    observed['disagreements'] = evaluate(row, observed)
    return observed


def required_keys(capability):
    method = capability.rsplit('.', 1)[-1]
    return {row['key'] for row in cases() if row['method'] == method}


def validate(output):
    import traceback
    from PyFLASH._persistence import write_json
    from reproducibility.normality_association_validation import comparison
    original.check_parameter_coverage()
    paired, comparisons = [], []
    for row in cases():
        observed, error = None, None
        try:
            observed = observe(row, output)
        except Exception:
            error = traceback.format_exc()
        agrees = error is None and not observed['disagreements']
        comparisons.append(comparison(row['key'], 'predeclared-literal-workbook-client/native-openpyxl',
            'functional-full-workbook-and-report-transport/no-additional-estimator-vote',
            {'complete_workbook_and_report_matches': agrees}, {'complete_workbook_and_report_matches': True},
            agrees, True, 'exact-all-sheets-all-cells-merged-ranges-literal-types-and-report-text', 0., 0.))
        paired.append(dict(case=row, observed=observed, error=error, agrees=agrees))
    write_json(Path(output) / 'excel-workbook-schema-paired.json', dict(schema_version=1,
        complete_contract=False, scope=__doc__, covered_parameters=sorted(original.check_parameter_coverage()), rows=paired))
    return comparisons
